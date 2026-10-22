"""Build the ATS-friendly resume.

Usage:
  python3 build_resume.py Harrison-Pielke-Lombardo-Resume.docx
  python3 build_resume.py Harrison-Pielke-Lombardo-Resume-Contract.docx contract
Then convert to PDF: soffice --headless --convert-to pdf <file>.docx
Requires python-docx.
"""
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.55)
sec.top_margin = sec.bottom_margin = Inches(0.45)
RIGHT = sec.page_width - sec.left_margin - sec.right_margin

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0

props = doc.core_properties
props.title = "Harrison Pielke-Lombardo - Resume"
props.author = "Harrison Pielke-Lombardo"


def para(space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return r


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3A5F")):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)


def heading(text):
    p = para(6, 2)
    run(p, text.upper(), bold=True, size=11, color=ACCENT)
    bottom_border(p)


def line_with_date(left, right, bold=True, italic=False, space_before=0, sep="  |  "):
    # Inline, no tab stops: right-aligned tabs make PDF parsers detach the date from its line.
    p = para(space_before, 0)
    run(p, left, bold=bold, italic=italic)
    run(p, sep + right, bold=False, italic=italic)
    return p


def bullet(text):
    p = para(0, 1)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.18))
    run(p, "\u2022\t" + text)


def labeled(label, text):
    p = para(0, 1.5)
    run(p, label + ": ", bold=True)
    run(p, text)


CONTRACT = len(sys.argv) > 2 and sys.argv[2] == "contract"

# Header
p = para()
run(p, "Harrison Pielke-Lombardo", bold=True, size=20)
p = para()
if CONTRACT:
    run(p, "Senior Backend & Full-Stack Engineer  |  Available for Contract & 1099 Work", bold=True, size=12,
        color=ACCENT)
else:
    run(p, "Senior Backend & Full-Stack Software Engineer", bold=True, size=12, color=ACCENT)
p = para(2, 0)
run(p, "Denver, CO (remote or hybrid)  |  720-209-6249  |  pielkekid@gmail.com  |  "
       "linkedin.com/in/tuh8888  |  github.com/tuh8888", size=10)

heading("Summary")
p = para()
if CONTRACT:
    run(p, "Senior backend and full-stack engineer available for contract and 1099 engagements, remote or "
           "on-site in Colorado. 10 years of building production systems in Python, Clojure, TypeScript, and "
           "Java. Built the pipeline, rules engine, and AWS infrastructure behind a B2B SaaS platform handling "
           "1M+ documents per month, from zero to $2.1M ARR through acquisition. Now building Well Met, a web "
           "and mobile app, with AI coding agents. Owns projects end to end and works directly with clients.")
else:
    run(p, "Senior software engineer with 10 years of experience building production backend systems, "
           "data pipelines, and web and mobile applications in Python, Clojure, TypeScript, and Java. Core "
           "engineer on a 4-6 person team at a B2B SaaS logistics startup, building the pipeline that processed "
           "1M+ documents per month and helping grow the product from zero to $2.1M ARR through acquisition. "
           "Now building Well Met, a web and mobile dating app, solo with AI coding agents. Owns features end "
           "to end and has mentored, interviewed, and led client calls.")

if CONTRACT:
    heading("What I Can Deliver")
    labeled("Backend APIs and data pipelines", "REST and GraphQL services, high-volume document and data "
                                               "processing, rules engines")
    labeled("Database migrations", "zero-downtime production migrations, including Datomic to PostgreSQL "
                                   "at Datomic's size limits")
    labeled("AWS infrastructure", "Terraform-managed, event-driven systems on Lambda, ECS, SQS, SNS, and RDS")
    labeled("LLM and ML automation", "production document processing and verification that replaces "
                                     "manual review work")
    labeled("Clojure and ClojureScript", "new builds, web and mobile apps, and modernizing existing systems")

heading("Skills")
labeled("Languages", "Python, Clojure, ClojureScript, TypeScript, Java, JavaScript, SQL, Bash")
if not CONTRACT:
    labeled("Backend", "REST APIs, GraphQL, data pipelines, rules engines, event-driven architecture, "
                       "database migrations")
labeled("Frontend & Mobile", "TypeScript, ClojureScript, React (Reagent / re-frame), React Native (Expo), "
                             "JavaScript, WebSockets")
labeled("Data", "PostgreSQL, Datomic, DynamoDB, Redis, SQL, knowledge graphs (RDF)")
labeled("Cloud & DevOps", "AWS (Lambda, ECS, EC2, SQS, SNS, RDS, DynamoDB, CloudFront), Terraform, "
                          "Docker, GitHub Actions, CI/CD")
labeled("AI & ML", "LLM integration, PyTorch, NLP, agentic development with Claude Code")

heading("Experience")
line_with_date("Founder & Software Engineer", "Aug 2026 - Present")
line_with_date("Well Met (pre-launch dating app, getwellmet.app)", "Denver, CO", sep=", ", bold=False, italic=True)
bullet("Building a web, iOS, and Android dating app centered on structured, mutual-consent post-date reviews.")
bullet("Architected a Clojure backend on PostgreSQL with real-time WebSocket updates, and one shared "
       "ClojureScript / re-frame core for the web app, admin tool, and React Native (Expo) mobile app.")
bullet("Implemented ID verification (Persona), geohash-only location storage, and audit-logged photo and "
       "video moderation.")
bullet("Led development with AI coding agents (Claude Code) under written architecture guardrails and "
       "review, merging 450+ pull requests in two months with Playwright, unit, and CI test suites.")
if not CONTRACT:
    bullet("Wrote the AWS deployment (ECS Fargate, ALB with WAF, RDS, GuardDuty) as Terraform with GitHub "
           "Actions deploy workflows.")

line_with_date("Senior Software Engineer", "Sep 2022 - Mar 2026", space_before=4)
line_with_date("Peruse Technology Inc. (acquired)", "Remote", sep=", ", bold=False, italic=True)
bullet("Helped scale a B2B SaaS logistics platform from zero to $2.1M ARR and 70+ customers on a 4-6 person "
       "team through acquisition.")
bullet("Designed and built the core load processing pipeline and rules engine in Clojure and Python, "
       "handling 1M+ documents per month via REST and GraphQL.")
bullet("Trained and deployed LLM-based models in PyTorch to automate document processing and freight broker "
       "verification, replacing work that took customers entire back-office audit teams.")
bullet("Led a zero-downtime migration of the production database from Datomic to PostgreSQL as it neared "
       "Datomic's size limits.")
bullet("Shared ownership of event-driven AWS infrastructure as code in Terraform (Lambda, ECS, SQS, SNS, RDS); "
       "took on-call rotations and contributed to the TypeScript frontend.")
if CONTRACT:
    bullet("Served as technical contact on client calls and owned relationships with several client accounts.")
else:
    bullet("Owned features end to end with the product team; mentored an engineer, interviewed candidates, "
           "reviewed code, and owned several client relationships as technical contact.")

line_with_date("Research Software Engineer (Graduate Researcher)", "Aug 2016 - Aug 2022", space_before=4)
line_with_date("University of Colorado Anschutz Medical Campus", "Aurora, CO", sep=", ", bold=False, italic=True)
bullet("Led software for research projects across 4+ teams, shipping 5 open-source packages, 4 datasets, and "
       "6 publications.")
bullet("Built Knowtator, a Java and ClojureScript / re-frame annotation web app (rewritten from Java Swing), "
       "used for 100k+ annotations by clinicians and researchers.")

heading("Education, Open Source & Speaking")
line_with_date("M.S., Computational Bioscience, University of Colorado Anschutz Medical Campus", "2022",
               bold=False)
line_with_date("B.S., Applied Mathematics (minors: Computer Science, Biochemistry), University of Colorado "
               "Boulder", "2016", bold=False)
labeled("Open source & speaking", "author of datascript-shortest-path; contributor to day8/re-frame-undo; "
                                  "invited talk on AI in transportation, IDEAL institute, Chicago, 2024")

doc.save(OUT)
